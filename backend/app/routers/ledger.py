"""工厂台账：按规格记录每日生产/出货/报废/退货/期初，含库存结存计算与修改留痕。"""

import json
from datetime import date, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.audit import log_action
from app.database import get_db
from app.deps import require_manager
from app.models import Document, LedgerEdit, LedgerEntry, PendingEntry, Product, User
from app.services.xlsx_out import append_row

router = APIRouter(prefix="/api/ledger", tags=["工厂台账"])

# 类型 → 中文名；IN 类增加库存，OUT 类减少库存
TYPE_LABELS = {
    "production": "生产入库",
    "shipment": "出货",
    "scrap": "报废",
    "return": "退货",
    "initial": "期初",
}
IN_TYPES = {"initial", "production", "return"}
OUT_TYPES = {"shipment", "scrap"}

EDITABLE_FIELDS = ("entry_date", "product_id", "entry_type", "quantity", "unit_price", "amount", "customer")


class LedgerEntryRequest(BaseModel):
    entry_date: str = Field(min_length=8, max_length=10)
    product_id: int
    entry_type: str
    quantity: float = Field(gt=0)
    unit_price: float | None = None
    amount: float | None = None
    customer: str | None = Field(default=None, max_length=100)
    note: str = ""


class LedgerEntryUpdateRequest(BaseModel):
    entry_date: str | None = Field(default=None, min_length=8, max_length=10)
    product_id: int | None = None
    entry_type: str | None = None
    quantity: float | None = Field(default=None, gt=0)
    unit_price: float | None = None
    amount: float | None = None
    customer: str | None = Field(default=None, max_length=100)


class PendingResolveRequest(BaseModel):
    action: str  # create=新建产品并入库 / assign=认领已有产品 / ignore=忽略
    product_id: int | None = None
    unit: str | None = Field(default=None, max_length=20)
    category: str | None = Field(default=None, max_length=50)
    warn_stock: float | None = None


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError:
        raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")


def _validate_type(entry_type: str) -> str:
    if entry_type not in TYPE_LABELS:
        raise HTTPException(status_code=400, detail=f"类型必须是 {'/'.join(TYPE_LABELS.keys())}")
    return entry_type


def _validate_product(db: Session, product_id: int) -> Product:
    product = db.query(Product).filter(Product.id == product_id).first()
    if not product:
        raise HTTPException(status_code=404, detail="产品不存在")
    if not product.is_active:
        raise HTTPException(status_code=400, detail=f"产品「{product.name} {product.spec}」已停用，请先启用")
    return product


def _calc_amount(quantity: float, unit_price: float | None, amount: float | None) -> float | None:
    if amount is not None:
        return round(amount, 2)
    if unit_price is not None:
        return round(quantity * unit_price, 2)
    return None


def _entry_out(entry: LedgerEntry, products: dict[int, Product], docs: dict[int, str]) -> dict:
    product = products.get(entry.product_id)
    return {
        "id": entry.id,
        "entry_date": entry.entry_date.isoformat(),
        "product_id": entry.product_id,
        "product_name": product.name if product else "（已删除）",
        "spec": product.spec if product else "",
        "unit": product.unit if product else "",
        "category": product.category if product else "",
        "entry_type": entry.entry_type,
        "entry_type_label": TYPE_LABELS.get(entry.entry_type, entry.entry_type),
        "quantity": entry.quantity,
        "unit_price": entry.unit_price,
        "amount": entry.amount,
        "customer": entry.customer,
        "source": entry.source,
        "confidence": entry.confidence,
        "evidence": entry.evidence,
        "document_name": docs.get(entry.document_id) if entry.document_id else None,
    }


@router.get("/inventory")
def inventory(
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    """库存结存：期初 + 生产入库 + 退货 − 出货 − 报废，按产品实时计算。"""
    rows = (
        db.query(LedgerEntry.product_id, LedgerEntry.entry_type, func.sum(LedgerEntry.quantity))
        .group_by(LedgerEntry.product_id, LedgerEntry.entry_type)
        .all()
    )
    sums = {(pid, etype): float(total or 0) for pid, etype, total in rows}
    rows = []
    for product in db.query(Product).order_by(Product.category, Product.name, Product.spec).all():
        qty = {t: float(sums.get((product.id, t), 0) or 0) for t in TYPE_LABELS}
        balance = qty["initial"] + qty["production"] + qty["return"] - qty["shipment"] - qty["scrap"]
        rows.append(
            {
                "product_id": product.id,
                "name": product.name,
                "spec": product.spec,
                "unit": product.unit,
                "category": product.category,
                "initial": qty["initial"],
                "production": qty["production"],
                "shipment": qty["shipment"],
                "scrap": qty["scrap"],
                "return": qty["return"],
                "balance": round(balance, 2),
                "warn_stock": product.warn_stock,
                "low_stock": product.warn_stock is not None and balance <= product.warn_stock,
                "is_active": product.is_active,
            }
        )
    return rows


def _report_range(start: str | None, end: str | None) -> tuple[date, date]:
    today = date.today()
    try:
        end_date = date.fromisoformat(end) if end else today
        start_date = date.fromisoformat(start) if start else today - timedelta(days=today.weekday())
    except ValueError:
        raise HTTPException(status_code=400, detail="日期格式应为 YYYY-MM-DD")
    if start_date > end_date:
        raise HTTPException(status_code=400, detail="开始日期不能晚于结束日期")
    return start_date, end_date


def _build_report(db: Session, start_date: date, end_date: date) -> dict:
    """日台账报表：每规格逐日结存流转 + 期间规格汇总。

    开仓（期初结存）= start 之前所有条目的净变化；期间行只列有变动的 (日期, 产品)。
    """
    products = {
        p.id: p for p in db.query(Product).order_by(Product.category, Product.name, Product.spec).all()
    }

    # 开仓：start 之前的净结存
    open_sums = (
        db.query(LedgerEntry.product_id, LedgerEntry.entry_type, func.sum(LedgerEntry.quantity))
        .filter(LedgerEntry.entry_date < start_date)
        .group_by(LedgerEntry.product_id, LedgerEntry.entry_type)
        .all()
    )
    balances = {pid: 0.0 for pid in products}
    for pid, etype, total in open_sums:
        sign = 1.0 if etype in IN_TYPES else -1.0
        balances[pid] = balances.get(pid, 0.0) + sign * float(total or 0)

    # 期间条目：按 (日期, 产品) 聚五类型 + 出货金额
    period_rows = (
        db.query(LedgerEntry)
        .filter(LedgerEntry.entry_date >= start_date, LedgerEntry.entry_date <= end_date)
        .order_by(LedgerEntry.entry_date, LedgerEntry.id)
        .all()
    )
    grid: dict[tuple[date, int], dict] = {}
    for r in period_rows:
        cell = grid.setdefault(
            (r.entry_date, r.product_id),
            {"initial": 0.0, "production": 0.0, "return": 0.0, "shipment": 0.0, "scrap": 0.0, "ship_amount": 0.0},
        )
        cell[r.entry_type] = cell.get(r.entry_type, 0.0) + (r.quantity or 0)
        if r.entry_type == "shipment":
            cell["ship_amount"] += r.amount if r.amount is not None else (r.unit_price or 0) * (r.quantity or 0)

    # 逐日滚动结存
    daily = []
    running = dict(balances)
    for (d, pid) in sorted(grid):
        cell = grid[(d, pid)]
        product = products.get(pid)
        opening = running.get(pid, 0.0)
        closing = opening + cell["initial"] + cell["production"] + cell["return"] - cell["shipment"] - cell["scrap"]
        running[pid] = closing
        daily.append(
            {
                "date": d.isoformat(),
                "product_id": pid,
                "name": product.name if product else "（已删除）",
                "spec": product.spec if product else "",
                "unit": product.unit if product else "",
                "opening": round(opening, 2),
                **{t: round(cell[t], 2) for t in ("initial", "production", "return", "shipment", "scrap")},
                "closing": round(closing, 2),
                "ship_amount": round(cell["ship_amount"], 2),
            }
        )

    # 期间规格汇总
    product_stats = []
    for pid, product in products.items():
        agg = {"initial": 0.0, "production": 0.0, "return": 0.0, "shipment": 0.0, "scrap": 0.0, "ship_amount": 0.0}
        for (d, gpid), cell in grid.items():
            if gpid == pid:
                for t in agg:
                    agg[t] += cell[t]
        opening = balances.get(pid, 0.0)
        closing = opening + agg["initial"] + agg["production"] + agg["return"] - agg["shipment"] - agg["scrap"]
        product_stats.append(
            {
                "product_id": pid,
                "name": product.name,
                "spec": product.spec,
                "unit": product.unit,
                "category": product.category,
                "opening": round(opening, 2),
                **{t: round(agg[t], 2) for t in ("initial", "production", "return", "shipment", "scrap")},
                "closing": round(closing, 2),
                "ship_amount": round(agg["ship_amount"], 2),
                "warn_stock": product.warn_stock,
            }
        )

    summary = {
        "production": round(sum(p["production"] for p in product_stats), 2),
        "shipment": round(sum(p["shipment"] for p in product_stats), 2),
        "scrap": round(sum(p["scrap"] for p in product_stats), 2),
        "return": round(sum(p["return"] for p in product_stats), 2),
        "ship_amount": round(sum(p["ship_amount"] for p in product_stats), 2),
        "row_count": len(daily),
    }
    return {
        "range": {"start": start_date.isoformat(), "end": end_date.isoformat()},
        "daily": daily,
        "products": product_stats,
        "summary": summary,
    }


@router.get("/report")
def ledger_report(
    start: str | None = Query(None),
    end: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    """日台账报表（每规格结存流转）。"""
    start_date, end_date = _report_range(start, end)
    return _build_report(db, start_date, end_date)


@router.get("/report/export")
def ledger_report_export(
    start: str | None = Query(None),
    end: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    """日台账报表导出 Excel：Sheet1 日流转，Sheet2 规格汇总。"""
    from io import BytesIO
    from urllib.parse import quote

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    start_date, end_date = _report_range(start, end)
    report = _build_report(db, start_date, end_date)

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="D9E8E4")
    bold = Font(bold=True)

    def _sheet(ws, title, headers, widths, rows):
        ws.title = title
        ws.append(headers)
        for cell in ws[1]:
            cell.font = bold
            cell.fill = header_fill
        for row in rows:
            append_row(ws, row)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    _sheet(
        wb.active, "日台账",
        ["日期", "品名", "规格", "单位", "开仓", "期初调整", "生产入库", "退货", "出货", "报废", "结存", "出货金额(元)"],
        [12, 14, 22, 8, 10, 10, 10, 10, 10, 10, 10, 12],
        [
            [d["date"], d["name"], d["spec"], d["unit"], d["opening"], d["initial"],
             d["production"], d["return"], d["shipment"], d["scrap"], d["closing"], d["ship_amount"]]
            for d in report["daily"]
        ],
    )
    _sheet(
        wb.create_sheet("规格汇总"),
        "规格汇总",
        ["品名", "规格", "单位", "分类", "期初结存", "期初调整", "生产入库", "退货", "出货", "报废", "期末结存", "出货金额(元)"],
        [14, 22, 8, 10, 10, 10, 10, 10, 10, 10, 10, 12],
        [
            [p["name"], p["spec"], p["unit"], p["category"], p["opening"], p["initial"],
             p["production"], p["return"], p["shipment"], p["scrap"], p["closing"], p["ship_amount"]]
            for p in report["products"]
        ],
    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = quote(f"日台账报表_{start_date.isoformat()}_{end_date.isoformat()}.xlsx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("/pending")
def list_pending(
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    """待确认条目：文档提取出但无法自动匹配产品的规格条目。"""
    rows = (
        db.query(PendingEntry)
        .order_by(PendingEntry.entry_date.desc(), PendingEntry.id.desc())
        .limit(500)
        .all()
    )
    doc_ids = {r.document_id for r in rows if r.document_id}
    docs = {d.id: d.filename for d in db.query(Document).filter(Document.id.in_(doc_ids)).all()} if doc_ids else {}
    items = []
    for r in rows:
        items.append(
            {
                "id": r.id,
                "document_id": r.document_id,
                "document_name": docs.get(r.document_id),
                "entry_date": r.entry_date.isoformat(),
                "product_name": r.product_name,
                "spec": r.spec,
                "entry_type": r.entry_type,
                "entry_type_label": TYPE_LABELS.get(r.entry_type, r.entry_type),
                "quantity": r.quantity,
                "unit": r.unit,
                "unit_price": r.unit_price,
                "amount": r.amount,
                "customer": r.customer,
                "source": r.source,
                "evidence": r.evidence,
                "created_at": r.created_at,
            }
        )
    return {"items": items, "total": len(items)}


@router.post("/pending/{pending_id}/resolve")
def resolve_pending(
    pending_id: int,
    body: PendingResolveRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_manager),
):
    """处理待确认条目：新建产品并入库 / 认领已有产品入库 / 忽略。"""
    pending = db.query(PendingEntry).filter(PendingEntry.id == pending_id).first()
    if not pending:
        raise HTTPException(status_code=404, detail="待确认条目不存在")

    name = (pending.product_name or "").strip()
    spec = (pending.spec or "").strip()
    desc = f"{pending.entry_date} {name} {spec} {TYPE_LABELS.get(pending.entry_type)} {pending.quantity}{pending.unit or ''}"

    if body.action == "ignore":
        db.delete(pending)
        db.commit()
        log_action(db, user, "忽略待确认条目", desc)
        return {"message": "已忽略"}

    if body.action == "assign":
        if not body.product_id:
            raise HTTPException(status_code=400, detail="认领时必须选择产品")
        product = _validate_product(db, body.product_id)
    elif body.action == "create":
        if not name and not spec:
            raise HTTPException(status_code=400, detail="品名与规格均为空，无法新建产品")
        dup = db.query(Product).filter(Product.name == (name or "未命名产品"), Product.spec == spec).first()
        if dup:
            raise HTTPException(status_code=400, detail=f"产品「{dup.name} {dup.spec}」已存在，请改用「认领产品」")
        product = Product(
            name=name or "未命名产品",
            spec=spec,
            unit=(body.unit or pending.unit or "千件").strip()[:20],
            category=(body.category or "螺母").strip()[:50],
            warn_stock=body.warn_stock,
        )
        db.add(product)
        db.flush()
    else:
        raise HTTPException(status_code=400, detail="action 必须是 create / assign / ignore")

    entry = LedgerEntry(
        document_id=pending.document_id,
        entry_date=pending.entry_date,
        product_id=product.id,
        entry_type=pending.entry_type,
        quantity=pending.quantity,
        unit_price=pending.unit_price,
        amount=pending.amount,
        customer=pending.customer,
        source=pending.source,
        confidence=pending.confidence,
        evidence=pending.evidence or "待确认条目入库",
    )
    db.add(entry)
    db.delete(pending)
    db.commit()
    log_action(db, user, "待确认条目入库", f"{desc} → {product.name} {product.spec}")
    return {"message": "已入库", "entry_id": entry.id}


@router.get("")
def list_entries(
    start: str | None = Query(None),
    end: str | None = Query(None),
    product_id: int | None = Query(None),
    entry_type: str | None = Query(None),
    customer: str | None = Query(None),
    source: str | None = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    """台账明细（SQL 分页），响应含筛选范围的服务端汇总（出货金额）。"""
    conds = []
    if start:
        conds.append(LedgerEntry.entry_date >= _parse_date(start))
    if end:
        conds.append(LedgerEntry.entry_date <= _parse_date(end))
    if product_id:
        conds.append(LedgerEntry.product_id == product_id)
    if entry_type:
        conds.append(LedgerEntry.entry_type == _validate_type(entry_type))
    if customer:
        conds.append(LedgerEntry.customer.like(f"%{customer.strip()}%"))
    if source:
        conds.append(LedgerEntry.source == source)

    total = db.query(LedgerEntry).filter(*conds).count()
    ship_amount = (
        db.query(func.coalesce(func.sum(LedgerEntry.amount), 0.0))
        .filter(*conds, LedgerEntry.entry_type == "shipment")
        .scalar()
    ) or 0.0

    rows = (
        db.query(LedgerEntry)
        .filter(*conds)
        .order_by(LedgerEntry.entry_date.desc(), LedgerEntry.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    product_ids = {r.product_id for r in rows}
    products = {p.id: p for p in db.query(Product).filter(Product.id.in_(product_ids)).all()} if product_ids else {}
    doc_ids = {r.document_id for r in rows if r.document_id}
    docs = {d.id: d.filename for d in db.query(Document).filter(Document.id.in_(doc_ids)).all()} if doc_ids else {}
    return {
        "items": [_entry_out(r, products, docs) for r in rows],
        "total": total,
        "page": page,
        "page_size": page_size,
        "summary": {"ship_amount": round(float(ship_amount), 2)},
    }


@router.post("", status_code=201)
def create_entry(
    body: LedgerEntryRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_manager),
):
    """手工台账录入（不依赖文档上传）。"""
    product = _validate_product(db, body.product_id)
    _validate_type(body.entry_type)
    entry = LedgerEntry(
        document_id=None,
        entry_date=_parse_date(body.entry_date),
        product_id=product.id,
        entry_type=body.entry_type,
        quantity=round(body.quantity, 4),
        unit_price=round(body.unit_price, 2) if body.unit_price is not None else None,
        amount=_calc_amount(body.quantity, body.unit_price, body.amount),
        customer=(body.customer or "").strip() or None,
        source="manual",
        confidence=1.0,
        evidence=f"手工台账：{body.note}" if body.note else "手工台账",
    )
    db.add(entry)
    db.commit()
    log_action(
        db, user, "手工台账",
        f"{entry.entry_date} {product.name} {product.spec} {TYPE_LABELS[entry.entry_type]} {entry.quantity}{product.unit}",
    )
    return {"id": entry.id, "message": "已入账"}


@router.put("/{entry_id}")
def update_entry(
    entry_id: int,
    body: LedgerEntryUpdateRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_manager),
):
    """修改台账，差异写入留痕表。"""
    entry = db.query(LedgerEntry).filter(LedgerEntry.id == entry_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="台账记录不存在")

    changes = []
    if body.entry_date is not None:
        new_date = _parse_date(body.entry_date)
        if new_date != entry.entry_date:
            changes.append({"field": "entry_date", "old": entry.entry_date.isoformat(), "new": new_date.isoformat()})
            entry.entry_date = new_date
    if body.product_id is not None and body.product_id != entry.product_id:
        product = _validate_product(db, body.product_id)
        changes.append({"field": "product_id", "old": entry.product_id, "new": product.id})
        entry.product_id = product.id
    if body.entry_type is not None and body.entry_type != entry.entry_type:
        _validate_type(body.entry_type)
        changes.append({"field": "entry_type", "old": entry.entry_type, "new": body.entry_type})
        entry.entry_type = body.entry_type
    for field in ("quantity", "unit_price", "amount", "customer"):
        new_value = getattr(body, field)
        if new_value is not None:
            old_value = getattr(entry, field)
            if old_value != new_value:
                changes.append({"field": field, "old": old_value, "new": new_value})
                setattr(entry, field, round(new_value, 4) if field == "quantity" else round(new_value, 2) if field != "customer" else new_value)

    if not changes:
        return {"message": "没有变化", "changes": []}

    # 金额未显式给出且有单价时，按 数量×单价 重算
    if body.amount is None and entry.unit_price is not None:
        entry.amount = round(entry.quantity * entry.unit_price, 2)

    db.add(
        LedgerEdit(
            entry_id=entry.id,
            user_id=user.id,
            username=user.username,
            changes=json.dumps(changes, ensure_ascii=False, default=str),
        )
    )
    db.commit()
    log_action(db, user, "修改台账", f"记录 #{entry.id}，{len(changes)} 处修改")
    return {"message": "已保存", "changes": changes}


@router.delete("/{entry_id}")
def delete_entry(
    entry_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_manager),
):
    entry = db.query(LedgerEntry).filter(LedgerEntry.id == entry_id).first()
    if not entry:
        raise HTTPException(status_code=404, detail="台账记录不存在")
    db.delete(entry)
    db.commit()
    log_action(db, user, "删除台账", f"记录 #{entry_id}（{entry.entry_date}）")
    return {"message": "已删除"}


@router.get("/{entry_id}/edits")
def entry_edits(
    entry_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    rows = (
        db.query(LedgerEdit)
        .filter(LedgerEdit.entry_id == entry_id)
        .order_by(LedgerEdit.id.desc())
        .all()
    )
    return [
        {
            "id": r.id,
            "username": r.username,
            "changes": json.loads(r.changes),
            "created_at": r.created_at,
        }
        for r in rows
    ]

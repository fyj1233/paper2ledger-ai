"""报表中心：日结单、分类排行、环比对比、Excel 导出。"""

import io
import json
from datetime import date, timedelta

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.deps import require_manager
from app.models import FinanceRecord, User
from app.services.xlsx_out import append_row

router = APIRouter(prefix="/api/reports", tags=["报表中心"])


def _range(start: str | None, end: str | None) -> tuple[date, date]:
    today = date.today()
    try:
        end_date = date.fromisoformat(end) if end else today
        start_date = date.fromisoformat(start) if start else end_date - timedelta(days=29)
    except ValueError:
        raise ValueError("日期格式应为 YYYY-MM-DD")
    return start_date, end_date


def _rows(db: Session, start: date, end: date, category: str | None):
    query = db.query(FinanceRecord).filter(
        FinanceRecord.biz_date >= start, FinanceRecord.biz_date <= end
    )
    if category:
        query = query.filter(FinanceRecord.category == category)
    return query.order_by(FinanceRecord.biz_date).all()


def _load_payments(raw: str | None) -> dict:
    if not raw:
        return {}
    try:
        return json.loads(raw) or {}
    except (ValueError, TypeError):
        return {}


def _daily_summary(rows):
    """按日聚合：金额、笔数、客单价、支付方式。"""
    daily: dict[date, dict] = {}
    for r in rows:
        d = daily.setdefault(
            r.biz_date,
            {"revenue": 0.0, "cost": 0.0, "profit": 0.0, "tx_count": 0, "records": 0, "payments": {}},
        )
        profit = r.profit if r.profit is not None else (r.revenue or 0) - (r.cost or 0)
        d["revenue"] += r.revenue or 0
        d["cost"] += r.cost or 0
        d["profit"] += profit
        d["records"] += 1
        d["tx_count"] += r.tx_count or 0
        for name, amount in _load_payments(r.payments).items():
            d["payments"][name] = d["payments"].get(name, 0) + amount

    result = []
    for d in sorted(daily):
        item = daily[d]
        tx = item["tx_count"]
        result.append({
            "date": d.isoformat(),
            "revenue": round(item["revenue"], 2),
            "cost": round(item["cost"], 2),
            "profit": round(item["profit"], 2),
            "tx_count": tx or None,
            "records": item["records"],
            # 有笔数才算客单价；否则留空避免误导
            "avg_ticket": round(item["revenue"] / tx, 2) if tx else None,
            "payments": {k: round(v, 2) for k, v in item["payments"].items()} or None,
        })
    return result


@router.get("/daily")
def daily_report(
    start: str | None = Query(None),
    end: str | None = Query(None),
    category: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    try:
        start_date, end_date = _range(start, end)
    except ValueError as e:
        return {"error": str(e)}
    rows = _rows(db, start_date, end_date, category)
    days = _daily_summary(rows)
    total = {
        "revenue": round(sum(d["revenue"] for d in days), 2),
        "cost": round(sum(d["cost"] for d in days), 2),
        "profit": round(sum(d["profit"] for d in days), 2),
        "tx_count": sum(d["tx_count"] or 0 for d in days) or None,
    }
    total["avg_ticket"] = (
        round(total["revenue"] / total["tx_count"], 2) if total["tx_count"] else None
    )
    return {"days": days, "total": total}


@router.get("/category")
def category_report(
    start: str | None = Query(None),
    end: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    try:
        start_date, end_date = _range(start, end)
    except ValueError as e:
        return {"error": str(e)}
    rows = _rows(db, start_date, end_date, None)
    grouped: dict[str, dict] = {}
    for r in rows:
        g = grouped.setdefault(r.category, {"revenue": 0.0, "cost": 0.0, "profit": 0.0, "records": 0})
        profit = r.profit if r.profit is not None else (r.revenue or 0) - (r.cost or 0)
        g["revenue"] += r.revenue or 0
        g["cost"] += r.cost or 0
        g["profit"] += profit
        g["records"] += 1
    items = [
        {"category": k, **{f: round(v[f], 2) for f in ("revenue", "cost", "profit")}, "records": v["records"]}
        for k, v in grouped.items()
    ]
    items.sort(key=lambda x: -x["revenue"])
    return {"items": items}


@router.get("/compare")
def compare_report(
    start: str | None = Query(None),
    end: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    try:
        start_date, end_date = _range(start, end)
    except ValueError as e:
        return {"error": str(e)}
    length = (end_date - start_date).days + 1
    prev_end = start_date - timedelta(days=1)
    prev_start = prev_end - timedelta(days=length - 1)

    def summary(rows):
        revenue = sum(r.revenue or 0 for r in rows)
        cost = sum(r.cost or 0 for r in rows)
        profit = sum(
            (r.profit if r.profit is not None else (r.revenue or 0) - (r.cost or 0)) for r in rows
        )
        return {"revenue": round(revenue, 2), "cost": round(cost, 2), "profit": round(profit, 2)}

    cur = summary(_rows(db, start_date, end_date, None))
    prev = summary(_rows(db, prev_start, prev_end, None))
    changes = {}
    for f in ("revenue", "cost", "profit"):
        if prev[f]:
            changes[f] = round((cur[f] - prev[f]) / prev[f], 4)
        else:
            changes[f] = None
    return {
        "current": {"start": start_date.isoformat(), "end": end_date.isoformat(), **cur},
        "previous": {"start": prev_start.isoformat(), "end": prev_end.isoformat(), **prev},
        "changes": changes,
    }


@router.get("/export")
def export_excel(
    start: str | None = Query(None),
    end: str | None = Query(None),
    category: str | None = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_manager),
):
    try:
        start_date, end_date = _range(start, end)
    except ValueError as e:
        return {"error": str(e)}

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = _rows(db, start_date, end_date, category)
    days = _daily_summary(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "日结单"
    header_fill = PatternFill("solid", fgColor="D9E8E4")
    bold = Font(bold=True)
    headers = ["日期", "营业额", "成本", "利润", "交易笔数", "记录数", "客单价", "支付方式"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = header_fill
    for d in days:
        append_row(ws, [
            d["date"], d["revenue"], d["cost"], d["profit"],
            d["tx_count"] or "", d["records"], d["avg_ticket"] or "",
            "；".join(f"{k} {v:.0f}元" for k, v in (d["payments"] or {}).items()),
        ])
    total_revenue = sum(d["revenue"] for d in days)
    total_cost = sum(d["cost"] for d in days)
    total_profit = sum(d["profit"] for d in days)
    ws.append(["合计", round(total_revenue, 2), round(total_cost, 2), round(total_profit, 2), "", "", "", ""])
    for cell in ws[ws.max_row]:
        cell.font = bold
    widths = [14, 12, 12, 12, 10, 8, 10, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 第二个 sheet：明细
    ws2 = wb.create_sheet("明细记录")
    headers2 = ["日期", "分类", "营业额", "成本", "利润", "笔数", "来源", "依据"]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = bold
        cell.fill = header_fill
    for r in rows:
        append_row(ws2, [
            r.biz_date.isoformat(), r.category, r.revenue, r.cost, r.profit,
            r.tx_count or "", r.source, (r.evidence or "")[:80],
        ])
    for i, w in enumerate([14, 10, 12, 12, 12, 8, 10, 50], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    from urllib.parse import quote

    filename = quote(f"经营报表_{start_date.isoformat()}_{end_date.isoformat()}.xlsx")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )

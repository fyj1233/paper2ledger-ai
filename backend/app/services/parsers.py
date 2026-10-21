"""文档解析：把上传的文件统一转成纯文本。"""

import csv
import io
import time

from app.config import settings

SUPPORTED_EXTS = {".txt", ".csv", ".docx", ".xlsx", ".pdf", ".jpg", ".jpeg", ".png"}

FORMAT_NAMES = {
    ".txt": "文本文档",
    ".csv": "CSV 表格",
    ".docx": "Word 文档",
    ".xlsx": "Excel 表格",
    ".pdf": "PDF 文档",
    ".jpg": "图片(OCR)",
    ".jpeg": "图片(OCR)",
    ".png": "图片(OCR)",
}

IMAGE_EXTS = {".jpg", ".jpeg", ".png"}


class ParseError(Exception):
    pass


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gbk", "utf-16"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def _parse_txt(data: bytes) -> str:
    return _decode_text(data)


def _parse_csv(data: bytes) -> str:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))
    lines = []
    for row in rows:
        cells = [c.strip() for c in row if c.strip()]
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _parse_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _parse_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def _parse_pdf(data: bytes) -> str:
    import pdfplumber

    texts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            texts.append(page.extract_text() or "")
    text = "\n".join(t for t in texts if t.strip())
    if len(text.strip()) < 10:
        raise ParseError("PDF 中未提取到文字，可能是扫描件/图片型 PDF，暂不支持，请上传文字版或改用 Word/Excel")
    return text


_ocr_instance = None


def _parse_image(data: bytes) -> str:
    """图片 OCR：本地离线识别（RapidOCR），适合小票/日报拍照。"""
    global _ocr_instance
    try:
        import cv2
        import numpy as np
    except ImportError:
        raise ParseError("服务端未安装 OCR 依赖（opencv-python），请联系管理员")

    arr = np.frombuffer(data, np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise ParseError("图片文件无法解码，请换一张图片试试")

    started = time.time()
    try:
        if _ocr_instance is None:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_instance = RapidOCR()
        result, _ = _ocr_instance(img)
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"图片识别失败：{e}")
    took = time.time() - started
    if not result:
        raise ParseError("图片中未识别出文字，请拍摄清晰一些（光线充足、文字朝上）再试")
    lines = [item[1] for item in result if item and len(item) > 1 and item[1]]
    text = "\n".join(lines)
    if len(text.strip()) < 5:
        raise ParseError("图片中未识别出有效文字，请拍摄清晰一些再试")
    print(f"[OCR] 识别完成：{len(lines)} 行，耗时 {took:.1f}s")
    return text


_PARSERS = {
    ".txt": _parse_txt,
    ".csv": _parse_csv,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
    ".pdf": _parse_pdf,
    ".jpg": _parse_image,
    ".jpeg": _parse_image,
    ".png": _parse_image,
}


def parse_file(data: bytes, ext: str) -> str:
    ext = ext.lower()
    if not ext.startswith("."):
        ext = f".{ext}"
    if ext not in SUPPORTED_EXTS:
        raise ParseError(f"不支持的文件格式 {ext}，支持：{', '.join(sorted(SUPPORTED_EXTS))}")
    parser = _PARSERS[ext]
    text = parser(data)
    if len(text) > settings.MAX_TEXT_CHARS:
        text = text[: settings.MAX_TEXT_CHARS]
    if not text.strip():
        raise ParseError("文档内容为空，未提取到任何文字")
    return text
